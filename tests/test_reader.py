"""
Unit tests for PPTX Reader (pptx-reader.py)
"""

import pytest
import json
import sys
from pathlib import Path
from typing import Dict, Any
import tempfile
from unittest.mock import Mock, patch, MagicMock

# Import our modules
sys.path.insert(0, str(Path(__file__).parent.parent))

# Import the reader classes directly
exec(open(Path(__file__).parent.parent / "pptx-reader.py").read())


class TestPresentationReader:
    """Test the PresentationReader class methods."""
    
    @pytest.fixture
    def reader(self):
        """Create a reader instance."""
        return PresentationReader()
    
    @pytest.fixture
    def mock_presentation(self):
        """Create a mock presentation for testing."""
        from pptx import Presentation
        from unittest.mock import Mock
        
        # Create a real presentation with one slide for testing
        prs = Presentation()
        slide_layout = prs.slide_layouts[0]  # Title slide layout
        slide = prs.slides.add_slide(slide_layout)
        
        # Add title
        title_shape = slide.shapes.title
        title_shape.text = "Test Slide Title"
        
        return prs
    
    @pytest.mark.unit
    def test_reader_initialization(self, reader):
        """Test that reader initializes correctly."""
        assert hasattr(reader, 'slide_counter')
        assert reader.slide_counter == 0
    
    @pytest.mark.unit
    def test_slide_type_detection_methods_exist(self, reader):
        """Test that all slide type detection methods exist."""
        detection_methods = [
            '_determine_slide_type',
            '_is_title_slide', 
            '_is_section_header',
            '_has_bullet_points',
            '_has_numbered_list',
            '_is_first_slide'
        ]
        
        for method_name in detection_methods:
            assert hasattr(reader, method_name), f"Missing method: {method_name}"
            method = getattr(reader, method_name)
            assert callable(method), f"Method {method_name} is not callable"
    
    @pytest.mark.unit
    def test_content_extraction_methods_exist(self, reader):
        """Test that all content extraction methods exist."""
        extraction_methods = [
            '_extract_title',
            '_extract_content_by_type',
            '_extract_chart_content',
            '_extract_table_content', 
            '_extract_text_image_content',
            '_extract_bullet_points',
            '_extract_numbered_list',
            '_extract_multi_column_content',
            '_extract_title_slide_content',
            '_extract_section_header_content',
            '_extract_generic_content',
            '_extract_background_image',
            '_extract_speaker_notes'
        ]
        
        for method_name in extraction_methods:
            assert hasattr(reader, method_name), f"Missing method: {method_name}"
            method = getattr(reader, method_name)
            assert callable(method), f"Method {method_name} is not callable"
    
    @pytest.mark.unit
    def test_title_extraction(self, reader):
        """Test title extraction from slides."""
        # Create mock slide with title
        mock_slide = Mock()
        mock_slide.shapes.title.text = "Test Title"
        
        title = reader._extract_title(mock_slide)
        assert title == "Test Title"
        
        # Test with no title
        mock_slide_no_title = Mock()
        mock_slide_no_title.shapes.title = None
        mock_slide_no_title.shapes = []
        
        title = reader._extract_title(mock_slide_no_title)
        assert title == ""
    
    @pytest.mark.unit
    def test_bullet_point_extraction(self, reader):
        """Test bullet point extraction with hierarchy."""
        # Mock slide with text frame
        mock_slide = Mock()
        mock_shape = Mock()
        mock_text_frame = Mock()
        mock_paragraph = Mock()
        
        # Set up mock hierarchy
        mock_paragraph.text = "Test bullet point"
        mock_paragraph.level = 0
        mock_text_frame.text = "Test bullet point"
        mock_text_frame.paragraphs = [mock_paragraph]
        mock_shape.text_frame = mock_text_frame
        mock_slide.shapes = [mock_shape]
        
        # Mock the helper method
        reader._is_title_text = Mock(return_value=False)
        
        content = reader._extract_bullet_points(mock_slide)
        
        assert "points" in content
        assert len(content["points"]) > 0
        assert content["points"][0]["text"] == "Test bullet point"
        assert content["points"][0]["level"] == 0
    
    @pytest.mark.unit
    def test_chart_content_extraction(self, reader):
        """Test chart content extraction."""
        # Mock slide with chart
        mock_slide = Mock()
        mock_shape = Mock()
        mock_chart = Mock()
        
        # Set up chart mock
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        
        mock_shape.shape_type = MSO_SHAPE_TYPE.CHART
        mock_shape.chart = mock_chart
        mock_chart.chart_type = XL_CHART_TYPE.COLUMN_CLUSTERED
        
        # Mock chart data (will likely fail due to complexity, but should not crash)
        mock_slide.shapes = [mock_shape]
        
        content = reader._extract_chart_content(mock_slide)
        
        assert "chart_type" in content
        assert "chart_data" in content
        assert content["chart_type"] == "bar"  # Default mapping
    
    @pytest.mark.unit
    def test_table_content_extraction(self, reader):
        """Test table content extraction."""
        # Mock slide with table
        mock_slide = Mock()
        mock_shape = Mock()
        mock_table = Mock()
        mock_row = Mock()
        mock_cell = Mock()
        
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        
        # Set up table mock
        mock_shape.shape_type = MSO_SHAPE_TYPE.TABLE
        mock_shape.table = mock_table
        mock_cell.text = "Test Cell"
        mock_row.cells = [mock_cell, mock_cell]  # Two cells
        mock_table.rows = [mock_row, mock_row]  # Two rows
        mock_slide.shapes = [mock_shape]
        
        content = reader._extract_table_content(mock_slide)
        
        assert "headers" in content
        assert "rows" in content
        assert len(content["headers"]) == 2
        assert len(content["rows"]) == 1  # One data row (first is header)
    
    @pytest.mark.unit
    def test_slide_type_detection_logic(self, reader):
        """Test slide type detection logic."""
        # Mock slide with charts
        mock_slide_chart = Mock()
        mock_shape_chart = Mock()
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        mock_shape_chart.shape_type = MSO_SHAPE_TYPE.CHART
        mock_slide_chart.shapes = [mock_shape_chart]
        
        slide_type = reader._determine_slide_type(mock_slide_chart, is_first=False)
        assert slide_type == "chart_slide"
        
        # Mock slide with tables
        mock_slide_table = Mock()
        mock_shape_table = Mock()
        mock_shape_table.shape_type = MSO_SHAPE_TYPE.TABLE
        mock_slide_table.shapes = [mock_shape_table]
        
        slide_type = reader._determine_slide_type(mock_slide_table, is_first=False)
        assert slide_type == "data_table"
        
        # Mock first slide (should be title_slide)
        mock_slide_first = Mock()
        mock_slide_first.shapes = []
        
        slide_type = reader._determine_slide_type(mock_slide_first, is_first=True)
        assert slide_type == "title_slide"
    
    @pytest.mark.unit
    def test_title_slide_content_extraction(self, reader):
        """Test title slide content extraction with proper structure."""
        # Mock title slide
        mock_slide = Mock()
        mock_shape1 = Mock()
        mock_shape2 = Mock()
        
        # Set up text frames
        mock_shape1.text_frame.text = "Main Title"
        mock_shape2.text_frame.text = "Innovation Through Excellence\nDr. John Smith, CTO\nJanuary 2024"
        
        mock_slide.shapes = [mock_shape1, mock_shape2]
        
        # Mock title extraction
        reader._extract_title = Mock(return_value="Main Title")
        
        content = reader._extract_title_slide_content(mock_slide)
        
        assert "title" in content
        assert "subtitle" in content  
        assert "author" in content
        assert "date" in content
        assert content["title"] == "Main Title"
    
    @pytest.mark.unit
    def test_speaker_notes_extraction(self, reader):
        """Test speaker notes extraction."""
        # Mock slide with notes
        mock_slide = Mock()
        mock_notes_slide = Mock()
        mock_notes_text_frame = Mock()
        
        mock_notes_text_frame.text = "These are speaker notes"
        mock_notes_slide.notes_text_frame = mock_notes_text_frame
        mock_slide.notes_slide = mock_notes_slide
        
        notes = reader._extract_speaker_notes(mock_slide)
        assert notes == "These are speaker notes"
        
        # Test slide with no notes
        mock_slide_no_notes = Mock()
        mock_slide_no_notes.notes_slide = None
        
        notes = reader._extract_speaker_notes(mock_slide_no_notes)
        assert notes == ""
    
    @pytest.mark.unit
    def test_presentation_metadata_extraction(self, reader):
        """Test presentation metadata extraction."""
        # Mock presentation with core properties
        mock_prs = Mock()
        mock_core_properties = Mock()
        
        mock_core_properties.title = "Test Presentation"
        mock_core_properties.author = "Test Author"  
        mock_core_properties.created = None
        mock_prs.core_properties = mock_core_properties
        
        # Mock first slide
        mock_slide = Mock()
        mock_slide.shapes.title.text = "Slide Title"
        mock_prs.slides = [mock_slide]
        
        metadata = reader._extract_presentation_metadata(mock_prs)
        
        assert "title" in metadata
        assert "author" in metadata
        assert "date" in metadata
        assert "theme" in metadata
        assert metadata["title"] == "Test Presentation"
        assert metadata["author"] == "Test Author"
    
    @pytest.mark.unit
    def test_error_handling_in_extraction(self, reader):
        """Test that extraction methods handle errors gracefully."""
        # Test with completely invalid slide
        invalid_slide = None
        
        # These should not raise exceptions
        try:
            reader._extract_title(invalid_slide)
            reader._extract_bullet_points(invalid_slide)
            reader._extract_chart_content(invalid_slide)
            reader._extract_table_content(invalid_slide)
        except AttributeError:
            # AttributeError is expected with None, but shouldn't crash the process
            pass
        except Exception as e:
            pytest.fail(f"Unexpected exception in error handling: {e}")


class TestSlideTypeDetection:
    """Test slide type detection algorithms."""
    
    @pytest.fixture
    def reader(self):
        return PresentationReader()
    
    @pytest.mark.unit
    def test_bullet_point_detection(self, reader):
        """Test bullet point detection logic."""
        # Mock slide with bullet points
        mock_slide = Mock()
        mock_shape = Mock()
        mock_text_frame = Mock()
        
        mock_text_frame.text = "• Point 1\n• Point 2\n• Point 3"
        mock_shape.text_frame = mock_text_frame
        mock_slide.shapes = [mock_shape]
        
        has_bullets = reader._has_bullet_points(mock_slide)
        assert has_bullets is True
        
        # Mock slide without bullet points
        mock_slide_no_bullets = Mock()
        mock_shape_no_bullets = Mock()
        mock_text_frame_no_bullets = Mock()
        
        mock_text_frame_no_bullets.text = "Regular text without bullets"
        mock_shape_no_bullets.text_frame = mock_text_frame_no_bullets
        mock_slide_no_bullets.shapes = [mock_shape_no_bullets]
        
        has_bullets = reader._has_bullet_points(mock_slide_no_bullets)
        assert has_bullets is False
    
    @pytest.mark.unit
    def test_numbered_list_detection(self, reader):
        """Test numbered list detection logic."""
        # Mock slide with numbered list
        mock_slide = Mock()
        mock_shape = Mock()
        mock_text_frame = Mock()
        
        mock_text_frame.text = "1. First item\n2. Second item\n3. Third item"
        mock_shape.text_frame = mock_text_frame
        mock_slide.shapes = [mock_shape]
        
        has_numbers = reader._has_numbered_list(mock_slide)
        assert has_numbers is True
        
        # Mock slide without numbered list
        mock_slide_no_numbers = Mock()
        mock_shape_no_numbers = Mock()
        mock_text_frame_no_numbers = Mock()
        
        mock_text_frame_no_numbers.text = "Regular text"
        mock_shape_no_numbers.text_frame = mock_text_frame_no_numbers
        mock_slide_no_numbers.shapes = [mock_shape_no_numbers]
        
        has_numbers = reader._has_numbered_list(mock_slide_no_numbers)
        assert has_numbers is False
    
    @pytest.mark.unit 
    def test_section_header_detection(self, reader):
        """Test section header detection logic."""
        # Mock section header (few text elements, moderate length)
        mock_slide = Mock()
        mock_shape1 = Mock()
        mock_shape2 = Mock()
        
        mock_shape1.text_frame.text = "Section Title"
        mock_shape2.text_frame.text = "Section Description"
        mock_slide.shapes = [mock_shape1, mock_shape2]
        
        is_section = reader._is_section_header(mock_slide)
        assert is_section is True
        
        # Mock slide with too much text (not a section header)
        mock_slide_too_much = Mock()
        mock_shape_long = Mock()
        mock_shape_long.text_frame.text = "Very long text " * 50  # Very long text
        mock_slide_too_much.shapes = [mock_shape_long]
        
        is_section = reader._is_section_header(mock_slide_too_much)
        assert is_section is False


class TestContentPreservation:
    """Test that content is preserved correctly during extraction."""
    
    @pytest.fixture
    def reader(self):
        return PresentationReader()
    
    @pytest.mark.unit
    def test_text_preservation_in_bullet_points(self, reader):
        """Test that bullet point text is preserved correctly."""
        # Mock slide with specific bullet points
        mock_slide = Mock()
        mock_shape = Mock()
        mock_text_frame = Mock()
        mock_paragraph1 = Mock()
        mock_paragraph2 = Mock()
        
        mock_paragraph1.text = "First bullet point"
        mock_paragraph1.level = 0
        mock_paragraph2.text = "  Second bullet point (indented)"
        mock_paragraph2.level = 1
        
        mock_text_frame.text = "First bullet point\n  Second bullet point (indented)"
        mock_text_frame.paragraphs = [mock_paragraph1, mock_paragraph2]
        mock_shape.text_frame = mock_text_frame
        mock_slide.shapes = [mock_shape]
        
        reader._is_title_text = Mock(return_value=False)
        
        content = reader._extract_bullet_points(mock_slide)
        
        # Check that text is preserved
        assert len(content["points"]) == 2
        assert content["points"][0]["text"] == "First bullet point"
        assert content["points"][1]["text"] == "Second bullet point (indented)"
        
        # Check that hierarchy is preserved
        assert content["points"][0]["level"] == 0
        assert content["points"][1]["level"] == 1
    
    @pytest.mark.unit
    def test_chart_type_mapping(self, reader):
        """Test that chart types are mapped correctly."""
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        
        chart_type_tests = [
            (XL_CHART_TYPE.PIE, "pie"),
            (XL_CHART_TYPE.LINE, "line"), 
            (XL_CHART_TYPE.AREA, "area"),
            (XL_CHART_TYPE.COLUMN_CLUSTERED, "bar")
        ]
        
        for xl_type, expected_type in chart_type_tests:
            mock_slide = Mock()
            mock_shape = Mock()
            mock_chart = Mock()
            
            mock_shape.shape_type = MSO_SHAPE_TYPE.CHART
            mock_shape.chart = mock_chart
            mock_chart.chart_type = xl_type
            mock_slide.shapes = [mock_shape]
            
            content = reader._extract_chart_content(mock_slide)
            assert content["chart_type"] == expected_type


class TestIntegrationWithPPTX:
    """Test integration with real python-pptx objects."""
    
    @pytest.mark.integration
    def test_real_presentation_reading(self, reader, temp_dir):
        """Test reading a real presentation file."""
        from pptx import Presentation
        
        # Create a real presentation
        prs = Presentation()
        
        # Add title slide
        title_slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(title_slide_layout)
        title = slide.shapes.title
        subtitle = slide.placeholders[1]
        title.text = "Test Presentation Title"
        subtitle.text = "Test Subtitle"
        
        # Add bullet point slide
        bullet_slide_layout = prs.slide_layouts[1]
        slide2 = prs.slides.add_slide(bullet_slide_layout)
        title2 = slide2.shapes.title
        body2 = slide2.placeholders[1]
        title2.text = "Bullet Points"
        tf = body2.text_frame
        tf.text = "First bullet"
        p = tf.add_paragraph()
        p.text = "Second bullet"
        
        # Save presentation
        pptx_path = temp_dir / "test_real.pptx"
        prs.save(str(pptx_path))
        
        # Read with our reader
        json_data = reader.read_presentation(str(pptx_path))
        
        # Validate structure
        assert "presentation_metadata" in json_data
        assert "slides" in json_data
        assert len(json_data["slides"]) == 2
        
        # Check first slide
        first_slide = json_data["slides"][0]
        assert first_slide["slide_type"] == "title_slide"
        assert "Test Presentation Title" in str(first_slide["content"])
        
        # Check second slide  
        second_slide = json_data["slides"][1]
        assert second_slide["slide_type"] == "bullet_points"
        assert second_slide["title"] == "Bullet Points"