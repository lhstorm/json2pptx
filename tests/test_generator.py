"""
Unit tests for PPTX Generator (pptx-generator.py)
"""

import pytest
import json
import sys
from pathlib import Path
from typing import Dict, Any
import tempfile
from unittest.mock import Mock, patch, MagicMock

# Import our modules
sys.path.insert(0, str(Path(__file__).parent.parent))

# Import the generator classes directly
exec(open(Path(__file__).parent.parent / "pptx-generator.py").read())


class TestPresentationGenerator:
    """Test the PresentationGenerator class methods."""
    
    @pytest.fixture
    def generator(self, sample_presentation_data):
        """Create a generator instance with sample data."""
        return PresentationGenerator(sample_presentation_data)
    
    @pytest.mark.unit
    def test_generator_initialization(self, sample_presentation_data):
        """Test that generator initializes correctly."""
        generator = PresentationGenerator(sample_presentation_data)
        
        assert generator.data == sample_presentation_data
        assert generator.theme is not None
        assert hasattr(generator, 'prs')
        assert len(generator.prs.slides) == 0  # No slides created yet
    
    @pytest.mark.unit
    def test_theme_extraction(self, sample_presentation_data):
        """Test theme extraction from metadata."""
        generator = PresentationGenerator(sample_presentation_data)
        
        expected_theme = sample_presentation_data["presentation_metadata"]["theme"]
        assert generator.theme["primary_color"] == expected_theme["primary_color"]
        assert generator.theme["secondary_color"] == expected_theme["secondary_color"]
        assert generator.theme["accent_color"] == expected_theme["accent_color"]
    
    @pytest.mark.unit
    def test_color_conversion(self, generator):
        """Test color conversion from hex to RGBColor."""
        from pptx.dml.color import RGBColor
        
        # Test valid hex colors
        color1 = generator._get_color("#FF0000")
        assert isinstance(color1, RGBColor)
        assert color1.rgb == RGBColor(255, 0, 0).rgb
        
        color2 = generator._get_color("#00FF00")
        assert color2.rgb == RGBColor(0, 255, 0).rgb
        
        # Test invalid color (should return default)
        color3 = generator._get_color("invalid")
        assert isinstance(color3, RGBColor)
    
    @pytest.mark.unit
    @pytest.mark.parametrize("slide_type", [
        "title_slide", "section_header", "bullet_points", "numbered_list",
        "chart_slide", "data_table", "comparison_table", "text_image_left",
        "text_image_right", "image_full", "content_single", "content_two_column",
        "content_three_column", "timeline", "process_flow", "quote_slide",
        "icon_points", "team_slide", "contact_slide", "thank_you"
    ])
    def test_slide_creation_methods_exist(self, generator, slide_type):
        """Test that all slide type creation methods exist."""
        method_name = f"_create_{slide_type}_slide"
        assert hasattr(generator, method_name), f"Missing method: {method_name}"
        
        method = getattr(generator, method_name)
        assert callable(method), f"Method {method_name} is not callable"
    
    @pytest.mark.unit
    def test_text_formatting(self, generator):
        """Test text formatting with different parameters."""
        from pptx.util import Pt
        from pptx.enum.text import PP_ALIGN
        
        # Create a mock text frame
        mock_slide = generator.prs.slides.add_slide(generator.prs.slide_layouts[0])
        mock_shape = mock_slide.shapes.add_textbox(
            PresentationConstants.MARGIN_SMALL, 
            PresentationConstants.MARGIN_SMALL,
            PresentationConstants.SLIDE_WIDTH - 2 * PresentationConstants.MARGIN_SMALL,
            PresentationConstants.TITLE_HEIGHT
        )
        text_frame = mock_shape.text_frame
        
        # Test basic text setting
        generator._set_text_with_formatting(
            text_frame, "Test Text", 18, alignment=PP_ALIGN.CENTER
        )
        
        assert text_frame.text == "Test Text"
        assert text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
        assert text_frame.paragraphs[0].runs[0].font.size == Pt(18)
    
    @pytest.mark.unit
    def test_chart_creation(self, generator):
        """Test chart creation methods."""
        slide = generator.prs.slides.add_slide(generator.prs.slide_layouts[0])
        
        chart_data = {
            "labels": ["A", "B", "C"],
            "datasets": [
                {"name": "Series 1", "values": [10, 20, 30]}
            ]
        }
        
        # Test chart creation doesn't raise errors
        try:
            generator._create_native_chart(
                slide, "bar", chart_data,
                PresentationConstants.MARGIN_LARGE,
                PresentationConstants.CONTENT_TOP,
                PresentationConstants.CHART_WIDTH,
                PresentationConstants.CHART_HEIGHT
            )
        except Exception as e:
            pytest.fail(f"Chart creation failed: {e}")
    
    @pytest.mark.unit
    def test_table_creation(self, generator):
        """Test table creation methods."""
        slide = generator.prs.slides.add_slide(generator.prs.slide_layouts[0])
        
        headers = ["Header 1", "Header 2", "Header 3"]
        rows = [
            ["Row 1 Col 1", "Row 1 Col 2", "Row 1 Col 3"],
            ["Row 2 Col 1", "Row 2 Col 2", "Row 2 Col 3"]
        ]
        
        # Test table creation doesn't raise errors
        try:
            generator._create_enhanced_table(
                slide, headers, rows,
                PresentationConstants.MARGIN_LARGE,
                PresentationConstants.CONTENT_TOP,
                PresentationConstants.SLIDE_WIDTH - 2 * PresentationConstants.MARGIN_LARGE,
                PresentationConstants.CONTENT_HEIGHT
            )
        except Exception as e:
            pytest.fail(f"Table creation failed: {e}")
    
    @pytest.mark.unit
    def test_image_download_and_add(self, generator):
        """Test image download and addition."""
        slide = generator.prs.slides.add_slide(generator.prs.slide_layouts[0])
        
        # Test with a real URL (should not actually download in unit test)
        with patch('requests.get') as mock_get:
            mock_response = Mock()
            mock_response.content = b"fake_image_data"
            mock_response.raise_for_status.return_value = None
            mock_get.return_value = mock_response
            
            # This should not raise an error
            generator._add_image_from_url(
                slide, "https://example.com/image.jpg",
                PresentationConstants.MARGIN_LARGE,
                PresentationConstants.CONTENT_TOP,
                PresentationConstants.CHART_WIDTH,
                PresentationConstants.CHART_HEIGHT
            )
    
    @pytest.mark.unit 
    def test_placeholder_shape_creation(self, generator):
        """Test placeholder shape creation."""
        slide = generator.prs.slides.add_slide(generator.prs.slide_layouts[0])
        
        generator._add_placeholder_shape(
            slide,
            PresentationConstants.MARGIN_LARGE,
            PresentationConstants.CONTENT_TOP,
            PresentationConstants.CHART_WIDTH,
            PresentationConstants.CHART_HEIGHT,
            "Test Placeholder"
        )
        
        # Check that a shape was added
        assert len(slide.shapes) > 0
    
    @pytest.mark.unit
    def test_error_handling_in_slide_creation(self, generator):
        """Test that slide creation handles errors gracefully."""
        
        # Test with invalid slide data
        invalid_slide_data = {
            "slide_number": 1,
            "slide_type": "invalid_type",
            "content": {}
        }
        
        # This should not raise an exception
        try:
            generator._create_slide(invalid_slide_data)
        except Exception as e:
            pytest.fail(f"Slide creation should handle invalid types gracefully: {e}")


class TestPresentationConstants:
    """Test the PresentationConstants class."""
    
    @pytest.mark.unit
    def test_constants_exist(self):
        """Test that all required constants exist."""
        required_constants = [
            'SLIDE_WIDTH', 'SLIDE_HEIGHT', 'MARGIN_SMALL', 'MARGIN_MEDIUM', 'MARGIN_LARGE',
            'FONT_TITLE', 'FONT_SUBTITLE', 'FONT_HEADING', 'FONT_BODY', 'FONT_CAPTION',
            'TITLE_TOP', 'TITLE_HEIGHT', 'CONTENT_TOP', 'CONTENT_HEIGHT',
            'CHART_WIDTH', 'CHART_HEIGHT'
        ]
        
        for constant in required_constants:
            assert hasattr(PresentationConstants, constant), f"Missing constant: {constant}"
    
    @pytest.mark.unit
    def test_constants_types(self):
        """Test that constants have correct types."""
        from pptx.util import Inches
        
        # Test dimension constants
        assert isinstance(PresentationConstants.SLIDE_WIDTH, type(Inches(1)))
        assert isinstance(PresentationConstants.SLIDE_HEIGHT, type(Inches(1)))
        
        # Test font size constants  
        assert isinstance(PresentationConstants.FONT_TITLE, int)
        assert isinstance(PresentationConstants.FONT_BODY, int)


class TestPresentationValidator:
    """Test the PresentationValidator class."""
    
    @pytest.mark.unit
    def test_valid_presentation_data(self, sample_presentation_data):
        """Test validation of valid presentation data."""
        try:
            # The validator is used internally, test that it doesn't reject valid data
            generator = PresentationGenerator(sample_presentation_data)
            assert generator.data == sample_presentation_data
        except Exception as e:
            pytest.fail(f"Valid presentation data was rejected: {e}")
    
    @pytest.mark.unit
    def test_invalid_presentation_data(self):
        """Test validation of invalid presentation data."""
        invalid_data = {
            "invalid_key": "invalid_value"
        }
        
        # Should either reject or handle gracefully
        try:
            generator = PresentationGenerator(invalid_data)
            # If it doesn't reject, it should at least handle gracefully
            assert hasattr(generator, 'data')
        except Exception:
            # Rejection is also acceptable
            pass


class TestIntegrationWithPPTX:
    """Test integration with python-pptx library."""
    
    @pytest.mark.integration
    def test_presentation_creation_with_all_slide_types(self, all_slide_types, temp_dir):
        """Test creating a presentation with all slide types."""
        
        # Create test data with all slide types
        slides_data = []
        for i, slide_type in enumerate(all_slide_types, 1):
            slide_data = {
                "slide_number": i,
                "slide_type": slide_type,
                "content": self._get_sample_content_for_type(slide_type),
                "speaker_notes": f"Notes for slide {i}",
                "transition": "fade",
                "duration": 15
            }
            slides_data.append(slide_data)
        
        presentation_data = {
            "presentation_metadata": {
                "title": "Test All Slide Types",
                "author": "Test Suite",
                "date": "2024-01-01",
                "theme": {
                    "primary_color": "#2C3E50",
                    "secondary_color": "#3498DB", 
                    "accent_color": "#E74C3C"
                }
            },
            "slides": slides_data
        }
        
        # Create generator and generate presentation
        generator = PresentationGenerator(presentation_data)
        output_path = temp_dir / "all_slide_types.pptx"
        
        try:
            generator.generate_presentation(str(output_path))
            assert output_path.exists()
            
            # Verify the PPTX can be opened
            from pptx import Presentation
            prs = Presentation(str(output_path))
            assert len(prs.slides) == len(all_slide_types)
            
        except Exception as e:
            pytest.fail(f"Failed to create presentation with all slide types: {e}")
    
    def _get_sample_content_for_type(self, slide_type: str) -> Dict[str, Any]:
        """Get sample content for each slide type."""
        
        content_templates = {
            "title_slide": {
                "title": "Test Title",
                "subtitle": "Test Subtitle", 
                "author": "Test Author",
                "date": "2024-01-01"
            },
            "section_header": {
                "title": "Section Header",
                "subtitle": "Section Description"
            },
            "bullet_points": {
                "points": [
                    {"text": "Point 1", "level": 0},
                    {"text": "Point 2", "level": 0}
                ]
            },
            "numbered_list": {
                "items": [
                    {"text": "Item 1", "level": 0},
                    {"text": "Item 2", "level": 0}
                ]
            },
            "chart_slide": {
                "chart_type": "bar",
                "chart_data": {
                    "labels": ["A", "B", "C"],
                    "datasets": [{"name": "Series 1", "values": [10, 20, 30]}]
                }
            },
            "data_table": {
                "headers": ["Col 1", "Col 2"],
                "rows": [["Data 1", "Data 2"]]
            },
            "comparison_table": {
                "items": ["Item 1", "Item 2"],
                "criteria": ["Criteria 1", "Criteria 2"]
            },
            "text_image_left": {
                "text": "Sample text",
                "image_path": "https://picsum.photos/400/300"
            },
            "text_image_right": {
                "text": "Sample text", 
                "image_path": "https://picsum.photos/400/300"
            },
            "image_full": {
                "image_path": "https://picsum.photos/800/600"
            },
            "content_single": {
                "text": "Single column content"
            },
            "content_two_column": {
                "columns": [
                    {"title": "Column 1", "content": "Content 1"},
                    {"title": "Column 2", "content": "Content 2"}
                ]
            },
            "content_three_column": {
                "columns": [
                    {"title": "Column 1", "content": "Content 1"},
                    {"title": "Column 2", "content": "Content 2"},
                    {"title": "Column 3", "content": "Content 3"}
                ]
            },
            "timeline": {
                "events": [
                    {"date": "2024-01", "title": "Event 1", "description": "Description 1"},
                    {"date": "2024-02", "title": "Event 2", "description": "Description 2"}
                ]
            },
            "process_flow": {
                "steps": [
                    {"title": "Step 1", "description": "Description 1"},
                    {"title": "Step 2", "description": "Description 2"}
                ]
            },
            "quote_slide": {
                "quote": "Test quote",
                "author": "Test Author",
                "title": "Quote Title"
            },
            "icon_points": {
                "points": [
                    {"icon": "⭐", "title": "Point 1", "description": "Description 1"}
                ]
            },
            "team_slide": {
                "members": [
                    {"name": "Person 1", "role": "Role 1", "image": "https://picsum.photos/200/200"}
                ]
            },
            "contact_slide": {
                "title": "Contact Us",
                "email": "test@example.com",
                "phone": "123-456-7890",
                "address": "Test Address"
            },
            "thank_you": {
                "title": "Thank You",
                "subtitle": "Questions?"
            }
        }
        
        return content_templates.get(slide_type, {"title": f"Test {slide_type}"})